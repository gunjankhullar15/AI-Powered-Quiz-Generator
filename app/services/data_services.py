import os
import PyPDF2
from docx import Document
from pptx import Presentation
from app.utils.splitter import get_splitter
from app.services.weaviate_services import transformer


def extract_text_from_pdf(file_path):
    """Extract text from PDF file."""
    with open(file_path, "rb") as f:
        reader = PyPDF2.PdfReader(f)
        return " ".join([page.extract_text() or "" for page in reader.pages])


def extract_text_from_docx(file_path):
    """Extract text from DOCX file."""
    doc = Document(file_path)
    # print(([para.text for para in doc.paragraphs]))
    return " ".join([para.text for para in doc.paragraphs])


def extract_text_from_txt(file_path):
    """Extract text from TXT file."""
    with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def extract_text_from_pptx(file_path):
    """Extract text from PPTX file."""
    prs = Presentation(file_path)
    text_runs = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text_runs.append(shape.text)
    return " ".join(text_runs)


def process_folder(path: str, client):
    """Process PDFs, DOCX, TXT, and PPTX files and store chunks + embeddings in Weaviate."""
    supported_ext = [".pdf", ".docx", ".txt", ".pptx"]
    files = os.listdir(path)
    supported_files = [f for f in files if os.path.splitext(f)[1].lower() in supported_ext]

    if not supported_files:
        return "No supported files (PDF, DOCX, TXT, PPTX) found in the folder."

    doc_collection = client.collections.get("Document")

    for file in supported_files:
        file_path = os.path.join(path, file)
        ext = os.path.splitext(file)[1].lower()

        # Extract text based on file type
        if ext == ".pdf":
            alltext = extract_text_from_pdf(file_path)
        elif ext == ".docx":
            alltext = extract_text_from_docx(file_path)
        elif ext == ".txt":
            alltext = extract_text_from_txt(file_path)
        elif ext == ".pptx":
            alltext = extract_text_from_pptx(file_path)
        else:
            continue  # Skip unsupported file types

        # Split text and create embeddings
        splitter = get_splitter()
        chunks = splitter.split_text(alltext)
        embeddings = transformer.encode(chunks).tolist()

        # Store in Weaviate
        for chunk, vector in zip(chunks, embeddings):
            doc_collection.data.insert(
                properties={"content": chunk, "filename": file},
                vector=vector
            )

    return f"✅ Successfully processed {len(supported_files)} files (PDF, DOCX, TXT, PPTX)."


